# -*- coding: utf-8 -*-
import openpyxl, re, sys, json, os, zipfile, xml.etree.ElementTree as ET

sys.stdout.reconfigure(encoding="utf-8")

FAUCET_ZH = {
    "stainless steel body": "不锈钢阀体",
    "ceramic cartridge": "陶瓷阀芯",
    "abs handle": "ABS把手",
    "alloy handle": "合金把手",
    "alloy body": "合金阀体",
    "stainless steel handle": "不锈钢把手",
    "hot and cold": "冷热水",
    "single cold": "单冷",
    "cold only": "单冷",
    "gun grey": "枪灰色",
    "brushed": "拉丝",
    "abs handwheel": "ABS手轮",
    "stainless steel handwheel": "不锈钢手轮",
    "abs hand": "ABS手轮",
    "stainless steel 201": "201不锈钢",
    "brass body": "铜阀体",
    "space aluminum body": "太空铝阀体",
    "space aluminum handle": "太空铝把手",
    "double rod": "双杆",
    "single rod": "单杆",
    "toilet brush": "马桶刷",
    "stainless steel fittings abs drain pipe": "不锈钢配件ABS排水管",
    "odor-proof": "防臭",
    "black tube": "黑色软管",
    "shower tube": "淋浴管",
    "abs shower head": "ABS花洒",
    "stainless steel tube": "不锈钢管",
    "instant electric water heater": "即热式电热水器",
}

def translate_faucet(line):
    if not line:
        return ""
    low = line.strip().lower()
    if "product size" in low:
        return line.replace("Product size", "产品规格").replace("Sigle", "单").replace("Double", "双")
    return FAUCET_ZH.get(low, line.strip())

def translate_desc(desc):
    if not desc:
        return ""
    parts = [p.strip() for p in desc.split("\n") if p.strip()]
    zh_parts = [translate_faucet(p) for p in parts]
    return "，".join(zh_parts)

def get_image_map(xlsx_path):
    mapping = {}
    with zipfile.ZipFile(xlsx_path, "r") as z:
        drawing_files = [f for f in z.namelist() if f.startswith("xl/drawings/drawing") and f.endswith(".xml")]
        for drawing_file in drawing_files:
            rels_file = drawing_file.replace(".xml", ".xml.rels")
            if rels_file not in z.namelist():
                basename = os.path.basename(drawing_file)
                rels_file = "xl/drawings/_rels/" + basename + ".rels"
            if rels_file not in z.namelist():
                continue
            rels_xml = z.read(rels_file).decode("utf-8")
            rid_to_file = {}
            root_rels = ET.fromstring(rels_xml)
            for rel in root_rels.findall("{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"):
                rid = rel.get("Id")
                target = rel.get("Target")
                if target.startswith("../"):
                    rid_to_file[rid] = target[3:]
                else:
                    rid_to_file[rid] = target
            drawing_xml = z.read(drawing_file).decode("utf-8")
            root = ET.fromstring(drawing_xml)
            ns = {
                "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
                "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            }
            for anchor in root.findall(".//xdr:twoCellAnchor", ns):
                frm = anchor.find("xdr:from", ns)
                col = int(frm.find("xdr:col", ns).text)
                row = int(frm.find("xdr:row", ns).text)
                blip = anchor.find(".//a:blip", ns)
                if blip is not None:
                    rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    if rid in rid_to_file:
                        img_file = rid_to_file[rid]
                        mapping[(row + 1, col + 1)] = img_file
    return mapping

def read_image(xlsx_path, image_path):
    if not image_path.startswith("xl/"):
        image_path = "xl/" + image_path
    with zipfile.ZipFile(xlsx_path, "r") as z:
        return z.read(image_path)

def get_ext(img_path):
    ext = os.path.splitext(img_path)[1].lstrip(".").lower()
    if ext in ("jpg", "jpeg"):
        return "jpeg"
    return "png"

for d in ["images/toilet", "images/faucet", "images/cabinet"]:
    os.makedirs(d, exist_ok=True)

# ===================== TOILET =====================
print("=== Toilet ===")
wb = openpyxl.load_workbook("D:/dabo/tutu马桶报价表.xlsx")
ws = wb[wb.sheetnames[0]]
img_map = get_image_map("D:/dabo/tutu马桶报价表.xlsx")

toilet_products = []
for row_idx in range(4, ws.max_row + 1):
    a_val = ws.cell(row=row_idx, column=1).value
    c_val = ws.cell(row=row_idx, column=3).value
    d_val = ws.cell(row=row_idx, column=4).value
    e_val = ws.cell(row=row_idx, column=5).value
    f_val = ws.cell(row=row_idx, column=6).value
    if not d_val and not c_val:
        continue
    seq = int(a_val) if a_val else None
    model = str(c_val).strip() if c_val else ""
    name = str(d_val).strip() if d_val else ""
    spec = str(e_val).strip() if e_val else ""
    price = int(f_val) if f_val else 0
    if seq:
        pid = "TU-" + str(seq)
    elif model:
        pid = "TU-" + model
    else:
        continue
    if "立柱盆" in name:
        ptype = "立柱盆"
    elif "踗便器" in name:
        ptype = "踗便器"
    else:
        ptype = "马桶"
    toilet_products.append({
        "id": pid, "model": model, "name": name, "spec": spec,
        "type": ptype, "price": price, "_row": row_idx
    })

for p in toilet_products:
    row = p["_row"]
    if (row, 2) in img_map:
        img_path = img_map[(row, 2)]
        try:
            img_data = read_image("D:/dabo/tutu马桶报价表.xlsx", img_path)
            ext = get_ext(img_path)
            filename = "images/toilet/toilet_" + p["id"] + "." + ext
            with open(filename, "wb") as f:
                f.write(img_data)
            p["image"] = filename
            print("  " + p["id"] + " -> " + filename + " (" + str(len(img_data)) + " bytes)")
        except Exception as e:
            print("  Error " + p["id"] + ": " + str(e))

for p in toilet_products:
    del p["_row"]

# ===================== FAUCET =====================
print("\n=== Faucet ===")
wb2 = openpyxl.load_workbook("D:/dabo/202603.xlsx")
ws2 = wb2[wb2.sheetnames[0]]
img_map2 = get_image_map("D:/dabo/202603.xlsx")

faucet_products = []
for row_idx in range(2, ws2.max_row + 1):
    b_val = ws2.cell(row=row_idx, column=2).value
    c_val = ws2.cell(row=row_idx, column=3).value
    e_val = ws2.cell(row=row_idx, column=5).value
    if not e_val:
        continue
    code = str(e_val).strip()
    desc_en = str(b_val).strip() if b_val else ""
    desc_zh = translate_desc(desc_en)
    price = int(c_val) if c_val else 0
    if code.startswith("HJ-1"): ptype = "角阀"
    elif code.startswith("HJ-2"): ptype = "淋浴套装"
    elif code.startswith("HJ-3"): ptype = "面盆龙头"
    elif code.startswith("HJ-4"): ptype = "面盆龙头"
    elif code.startswith("HJ-5"):
        if "5550" in code or "5551" in code: ptype = "淋浴管"
        elif "5552" in code: ptype = "马桶刷"
        elif "5553" in code or "5554" in code: ptype = "地漏"
        else: ptype = "面盆龙头"
    elif code.startswith("HJ-6"): ptype = "厨房龙头"
    elif code.startswith("HJ-7"):
        if "7701" in code: ptype = "淋浴管"
        else: ptype = "角阀"
    elif code.startswith("HJ-8"):
        if "862" in code: ptype = "淋浴套装"
        else: ptype = "地漏"
    elif code.startswith("HJ-9"): ptype = "地漏/配件"
    else: ptype = "龙头"
    name_zh = desc_zh if desc_zh else code
    faucet_products.append({
        "id": code, "code": code, "name": name_zh, "desc": desc_zh,
        "type": ptype, "price": price, "_row": row_idx
    })

for p in faucet_products:
    row = p["_row"]
    if (row, 1) in img_map2:
        img_path = img_map2[(row, 1)]
        try:
            img_data = read_image("D:/dabo/202603.xlsx", img_path)
            ext = get_ext(img_path)
            filename = "images/faucet/faucet_" + p["id"] + "." + ext
            with open(filename, "wb") as f:
                f.write(img_data)
            p["image"] = filename
            print("  " + p["id"] + " -> " + filename + " (" + str(len(img_data)) + " bytes)")
        except Exception as e:
            print("  Error " + p["id"] + ": " + str(e))

for p in faucet_products:
    del p["_row"]

# ===================== CABINET =====================
print("\n=== Cabinet ===")
wb3 = openpyxl.load_workbook("D:/dabo/浴室柜.xlsx")
ws3 = wb3[wb3.sheetnames[0]]
img_map3 = get_image_map("D:/dabo/浴室柜.xlsx")

cabinet_products = []
for row_idx in range(2, ws3.max_row + 1):
    b_val = ws3.cell(row=row_idx, column=2).value
    c_val = ws3.cell(row=row_idx, column=3).value
    d_val = ws3.cell(row=row_idx, column=4).value
    if not b_val:
        continue
    model = str(b_val).strip()
    if isinstance(c_val, (int, float)):
        price = int(c_val)
        spec = ""
    else:
        price = 0
        spec = str(c_val).strip() if c_val else ""
    note = str(d_val).strip() if d_val else ""
    if note:
        spec = (spec + " " + note).strip() if spec else note
    seq = row_idx - 1
    pid = "YC-" + str(seq)
    if "支架" in model: ptype = "支架"
    else: ptype = "浴室柜"
    cabinet_products.append({
        "id": pid, "model": model, "name": model, "spec": spec,
        "type": ptype, "price": price, "_row": row_idx
    })

for p in cabinet_products:
    row = p["_row"]
    if (row, 1) in img_map3:
        img_path = img_map3[(row, 1)]
        try:
            img_data = read_image("D:/dabo/浴室柜.xlsx", img_path)
            ext = get_ext(img_path)
            filename = "images/cabinet/cabinet_" + p["id"] + "." + ext
            with open(filename, "wb") as f:
                f.write(img_data)
            p["image"] = filename
            print("  " + p["id"] + " -> " + filename + " (" + str(len(img_data)) + " bytes)")
        except Exception as e:
            print("  Error " + p["id"] + ": " + str(e))

for p in cabinet_products:
    del p["_row"]

# ===================== SAVE =====================
products = {
    "toilet": toilet_products,
    "faucet": faucet_products,
    "cabinet": cabinet_products
}

with open("D:/dabo/products.json", "w", encoding="utf-8") as f:
    json.dump(products, f, ensure_ascii=False, indent=2)

with open("D:/dabo/products.js", "w", encoding="utf-8") as f:
    f.write("// Auto-generated by extract.py" + chr(10))
    f.write("const PRODUCTS = ")
    json.dump(products, f, ensure_ascii=False, indent=2)
    f.write(";" + chr(10))

print("\n=== Summary ===")
print("  Toilet: " + str(len(toilet_products)) + " products")
print("  Faucet: " + str(len(faucet_products)) + " products")
print("  Cabinet: " + str(len(cabinet_products)) + " products")
